import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from src.reference_style import extract_core_reference_style


class ReferenceStyleTests(unittest.TestCase):
    def test_palette_is_derived_from_reference_without_blue_fallback(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "red_reference.pptx"
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
            for index in range(3):
                shape = slide.shapes.add_shape(1, Inches(1 + index * 2), Inches(2), Inches(1.5), Inches(1))
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor(180, 32, 37)
            title = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8), Inches(0.6))
            run = title.text_frame.paragraphs[0].add_run()
            run.text = "红色参考模板"
            run.font.size = Pt(28)
            run.font.color.rgb = RGBColor(180, 32, 37)
            deck.save(path)

            style = extract_core_reference_style(path)
            self.assertEqual(style["palette"]["roles"]["primary"], "B42025")
            self.assertEqual(style["hard_constraints"]["policy"], "dynamic_from_core_reference")
            self.assertIn("B42025", style["hard_constraints"]["palette"]["allowed_hex"])
            self.assertNotIn("1B61A5", style["hard_constraints"]["palette"]["allowed_hex"])


if __name__ == "__main__":
    unittest.main()
